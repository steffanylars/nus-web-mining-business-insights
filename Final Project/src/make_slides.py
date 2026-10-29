"""Optional presentation generator using corrected measured artifacts only."""
import json
from config import RES,prepare

def main():
    from pptx import Presentation
    from pptx.util import Inches,Pt
    prepare()
    baseline=json.loads((RES/'results_baselines_lgb.json').read_text())
    neural=json.loads((RES/'results_neural.json').read_text())
    dataset=json.loads((RES/'dataset_manifest.json').read_text())
    uncertainty=json.loads((RES/'paired_bootstrap.json').read_text())
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    def slide(title,lines):
        s=prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text=title
        tf=s.placeholders[1].text_frame; tf.clear()
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line; p.font.size=Pt(23); p.space_after=Pt(16)
    slide('JDSearch corrected pipeline',[
        'Group 13 | Team review draft on ethan',
        'Personalised reranking of logged candidate slates',
        'Measured results replace the original draft figures',
        'Confirm course requirements and contribution statements before submission'])
    slide('What the history timeline shows',[
        'Reconstruct ages with suffix sums of consecutive time intervals',
        'Sharp click/cart drops around seven and 90 days',
        'Consistent with different logging windows; actual platform policy unverified',
        'Absent older clicks are unobserved, not negative preference evidence'])
    slide('Corrections to the pipeline',[
        'Parse the actual U+0018 term delimiter and retain category names',
        f'Preserve all {dataset["history_rows"]:,} sampled history events',
        'Sample the short and order horizons independently',
        'Fit shared preprocessing and item factors on training users only'])
    slide('Evaluation protocol',[
        'Fixed 20,000-user sample, split 70/15/15 by user',
        'Full logged slates at evaluation; 2,965 positive test queries',
        'NDCG gains [0,1,3,7]; binary relevance for MRR and precision',
        'Deterministic candidate-ID tie keys; no source-position feature'])
    values=[('Popularity',baseline['Popularity']),('SVD train-only',baseline['SVD(train-only fold-in)']),
        ('LambdaMART',baseline['LambdaMART(features)']),('MLP',neural['MLP(no towers)']),
        ('DHAR',neural['DHAR(ours)']),('History CF',neural['History-conditioned CF'])]
    slide('Measured test results',[f'{name}: NDCG@10 {m["NDCG@10"]:.4f}' for name,m in values])
    ci=uncertainty['dual_minus_mlp']
    slide('Does dual attention add value',[
        'Matched MLP, pooled, short-only, order-only and dual variants',
        f'DHAR minus MLP: {ci["delta"]:+.4f} NDCG@10',
        f'Paired query bootstrap 95% interval: [{ci["ci95"][0]:+.4f}, {ci["ci95"][1]:+.4f}]',
        'Use this comparison, not the tree comparison, to assess the added towers'])
    slide('Limits and team decisions',[
        'One training seed and previously inspected test partition',
        'Offline ranking does not establish causal conversion or revenue lift',
        'Full-slate subgroup diagnostics replace shortened warm/cold slates',
        'Discuss model complexity and final evaluation protocol with the team'])
    target=RES/'BT4222_Group13_Corrected_Presentation.pptx'
    prs.save(target); print(target)

if __name__=='__main__':
    main()
